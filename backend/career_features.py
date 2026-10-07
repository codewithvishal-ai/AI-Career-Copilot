import os
import re
from mongo_store import load_dataset


class CareerFeatures:

    # ============================================================
    # RESUME FILE PARSER
    # ============================================================

    def parse_resume_file(self, file_path):

        try:

            if not os.path.exists(file_path):

                print(
                    f"❌ File not found: {file_path}"
                )

                return ""

            extension = (
                os.path.splitext(file_path)[1]
                .lower()
            )

            # ----------------------------------------------------
            # TXT FILE
            # ----------------------------------------------------

            if extension == ".txt":

                with open(
                    file_path,
                    "r",
                    encoding="utf-8"
                ) as file:

                    return file.read()

            # ----------------------------------------------------
            # PDF FILE
            # ----------------------------------------------------

            elif extension == ".pdf":

                try:

                    import PyPDF2

                except ImportError:

                    print(
                        "\n❌ PyPDF2 is not installed."
                    )

                    print(
                        "Install it using:"
                    )

                    print(
                        "pip install PyPDF2"
                    )

                    return ""

                text = ""

                with open(
                    file_path,
                    "rb"
                ) as file:

                    reader = PyPDF2.PdfReader(
                        file
                    )

                    for page in reader.pages:

                        page_text = (
                            page.extract_text()
                        )

                        if page_text:

                            text += (
                                page_text
                                + "\n"
                            )

                return text

            # ----------------------------------------------------
            # DOCX FILE
            # ----------------------------------------------------

            elif extension == ".docx":

                try:

                    from docx import Document

                except ImportError:

                    print(
                        "\n❌ python-docx is not installed."
                    )

                    print(
                        "Install it using:"
                    )

                    print(
                        "pip install python-docx"
                    )

                    return ""

                document = Document(
                    file_path
                )

                text = ""

                for paragraph in document.paragraphs:

                    if paragraph.text.strip():

                        text += (
                            paragraph.text
                            + "\n"
                        )

                # Also read tables inside DOCX

                for table in document.tables:

                    for row in table.rows:

                        row_text = " ".join(
                            cell.text.strip()
                            for cell in row.cells
                        )

                        if row_text:

                            text += (
                                row_text
                                + "\n"
                            )

                return text

            # ----------------------------------------------------
            # UNSUPPORTED FILE
            # ----------------------------------------------------

            else:

                print(
                    "\n❌ Unsupported resume format."
                )

                print(
                    "Supported formats:"
                )

                print(
                    "PDF, DOCX, TXT"
                )

                return ""

        except Exception as e:

            print(
                f"\n❌ Error parsing resume: {e}"
            )

            return ""

    # ============================================================
    # EXTRACT SKILLS
    # ============================================================

    def extract_skills_from_resume(
        self,
        text
    ):

        known_skills = [

            "python",
            "java",
            "c",
            "c++",
            "c#",
            "javascript",
            "typescript",

            "html",
            "css",
            "react",
            "node",
            "node.js",
            "django",
            "flask",

            "sql",
            "mysql",
            "postgresql",
            "mongodb",

            "git",
            "github",
            "gitlab",

            "pandas",
            "numpy",
            "matplotlib",
            "seaborn",

            "machine learning",
            "deep learning",
            "artificial intelligence",
            "ai",

            "tensorflow",
            "pytorch",
            "scikit-learn",

            "data science",
            "data analysis",

            "excel",
            "power bi",

            "aws",
            "azure",
            "google cloud",

            "docker",
            "kubernetes",

            "linux",

            "oop",
            "object oriented programming",

            "rest api",
            "api",

            "communication",
            "leadership",
            "problem solving"
        ]

        text_lower = text.lower()

        detected_skills = []

        for skill in known_skills:

            if skill.lower() in text_lower:

                detected_skills.append(
                    skill
                )

        return detected_skills

    # ============================================================
    # EXTRACT EMAIL
    # ============================================================

    def extract_email(
        self,
        text
    ):

        pattern = (
            r"[A-Za-z0-9._%+-]+"
            r"@[A-Za-z0-9.-]+"
            r"\.[A-Za-z]{2,}"
        )

        match = re.search(
            pattern,
            text
        )

        if match:

            return match.group(0)

        return "Not detected"

    # ============================================================
    # EXTRACT PHONE
    # ============================================================

    def extract_phone(
        self,
        text
    ):

        patterns = [

            r"\+91[\s-]?[6-9]\d{9}",

            r"\b[6-9]\d{9}\b",

            r"\+91[\s-]?\d{10}"
        ]

        for pattern in patterns:

            match = re.search(
                pattern,
                text
            )

            if match:

                return match.group(0)

        return "Not detected"

    # ============================================================
    # EXTRACT NAME
    # ============================================================

    def extract_name(
        self,
        text
    ):

        lines = [

            line.strip()

            for line in text.splitlines()

            if line.strip()
        ]

        if not lines:

            return "Not detected"

        # Usually the first meaningful
        # line of a resume is the name.

        for line in lines[:10]:

            if (
                len(line.split()) <= 5
                and
                not re.search(
                    r"@|www\.|http|phone|mobile|"
                    r"resume|curriculum|linkedin|github",
                    line.lower()
                )
            ):

                # Avoid section headings

                ignored = [

                    "resume",
                    "curriculum vitae",
                    "cv",
                    "education",
                    "skills",
                    "projects",
                    "experience",
                    "certifications",
                    "summary",
                    "objective"
                ]

                if line.lower() not in ignored:

                    return line

        return "Not detected"

    # ============================================================
    # DETECT RESUME SECTIONS
    # ============================================================

    def detect_resume_sections(
        self,
        text
    ):

        text_lower = text.lower()

        sections = {

            "education": [

                "education",
                "academic",
                "qualification",
                "degree",
                "bca",
                "b.tech",
                "btech",
                "mca",
                "m.tech",
                "college",
                "university"
            ],

            "experience": [

                "experience",
                "work experience",
                "employment",
                "internship",
                "intern",
                "professional experience"
            ],

            "projects": [

                "projects",
                "project",
                "portfolio"
            ],

            "skills": [

                "skills",
                "technical skills",
                "technologies",
                "technical knowledge",
                "technical expertise"
            ],

            "certifications": [

                "certifications",
                "certificates",
                "certification",
                "credential"
            ],

            "summary": [

                "summary",
                "profile summary",
                "professional summary",
                "objective",
                "career objective"
            ]
        }

        detected = {}

        for section, keywords in sections.items():

            detected[section] = any(

                keyword in text_lower

                for keyword in keywords
            )

        return detected

    # ============================================================
    # RESUME ANALYSIS
    # ============================================================

    def resume_analysis(
        self
    ):

        print(
            "\n========== RESUME ANALYSIS =========="
        )

        print(
            "\nChoose Resume Input:"
        )

        print(
            "1. Upload/Parse Resume File"
        )

        print(
            "2. Enter Skills Manually"
        )

        choice = input(
            "\nEnter your choice: "
        ).strip()

        # ========================================================
        # FILE PARSING
        # ========================================================

        if choice == "1":

            file_path = input(
                "\nEnter resume file path: "
            ).strip()

            # Remove quotes if user pasted
            # a Windows path with quotes.

            file_path = file_path.strip(
                '"'
            )

            print(
                "\n📄 Reading resume..."
            )

            resume_text = (
                self.parse_resume_file(
                    file_path
                )
            )

            if not resume_text:

                print(
                    "\n❌ Could not extract text "
                    "from the resume."
                )

                return

            print(
                "✅ Resume successfully parsed!"
            )

            print(
                f"📄 Extracted characters: "
                f"{len(resume_text)}"
            )

        # ========================================================
        # MANUAL SKILLS
        # ========================================================

        elif choice == "2":

            skills = input(
                "\nEnter your skills "
                "(comma separated): "
            )

            resume_text = skills

        else:

            print(
                "\n❌ Invalid choice."
            )

            return

        # ========================================================
        # EXTRACT BASIC INFORMATION
        # ========================================================

        name = self.extract_name(
            resume_text
        )

        email = self.extract_email(
            resume_text
        )

        phone = self.extract_phone(
            resume_text
        )

        # ========================================================
        # EXTRACT SKILLS
        # ========================================================

        skills_list = (
            self.extract_skills_from_resume(
                resume_text
            )
        )

        # Remove duplicates
        # while preserving order.

        skills_list = list(
            dict.fromkeys(
                skills_list
            )
        )

        # ========================================================
        # DETECT SECTIONS
        # ========================================================

        sections = (
            self.detect_resume_sections(
                resume_text
            )
        )

        # ========================================================
        # RESUME REPORT
        # ========================================================

        print(
            "\n\n========== RESUME REPORT =========="
        )

        # ========================================================
        # PERSONAL INFORMATION
        # ========================================================

        print(
            "\n👤 PERSONAL INFORMATION"
        )

        print(
            f"Name  : {name}"
        )

        print(
            f"Email : {email}"
        )

        print(
            f"Phone : {phone}"
        )

        # ========================================================
        # SKILLS
        # ========================================================

        print(
            "\n🛠️ SKILLS"
        )

        if skills_list:

            for skill in skills_list:

                print(
                    f"✅ {skill.title()}"
                )

            print(
                f"\nTotal Skills: "
                f"{len(skills_list)}"
            )

        else:

            print(
                "❌ No known skills detected."
            )

        # ========================================================
        # RESUME SECTIONS
        # ========================================================

        print(
            "\n📑 RESUME SECTIONS"
        )

        for section, found in sections.items():

            section_name = (
                section.title()
            )

            if found:

                print(
                    f"✅ {section_name}"
                )

            else:

                print(
                    f"❌ {section_name} - Missing"
                )

        # ========================================================
        # REQUIRED SKILL CHECK
        # ========================================================

        required_skills = [

            "python",
            "git",
            "sql"
        ]

        print(
            "\n🔍 REQUIRED SKILL CHECK"
        )

        for skill in required_skills:

            if skill in skills_list:

                print(
                    f"✅ {skill.title()} - Available"
                )

            else:

                print(
                    f"❌ {skill.title()} - Missing"
                )

        # ========================================================
        # RESUME STRENGTH SCORE
        # ========================================================

        print(
            "\n💪 RESUME STRENGTH"
        )

        strength_score = 0

        # Skills

        if len(skills_list) >= 5:

            strength_score += 25

        elif len(skills_list) >= 3:

            strength_score += 15

        # Education

        if sections["education"]:

            strength_score += 20

        # Projects

        if sections["projects"]:

            strength_score += 20

        # Experience

        if sections["experience"]:

            strength_score += 20

        # Certifications

        if sections["certifications"]:

            strength_score += 15

        # Keep score at maximum 100.

        strength_score = min(
            strength_score,
            100
        )

        print(
            f"Resume Strength Score: "
            f"{strength_score}/100"
        )

        if strength_score >= 80:

            print(
                "🔥 Strong Resume Structure"
            )

        elif strength_score >= 60:

            print(
                "✅ Good Resume Structure"
            )

        elif strength_score >= 40:

            print(
                "👍 Average Resume Structure"
            )

        else:

            print(
                "⚠ Resume needs improvement"
            )

        # ========================================================
        # RECOMMENDATIONS
        # ========================================================

        print(
            "\n💡 RECOMMENDATIONS"
        )

        recommendations = []

        if len(skills_list) < 5:

            recommendations.append(
                "Add more relevant technical skills."
            )

        if not sections["projects"]:

            recommendations.append(
                "Add a Projects section."
            )

        if not sections["experience"]:

            recommendations.append(
                "Add internship or experience details."
            )

        if not sections["certifications"]:

            recommendations.append(
                "Add relevant certifications."
            )

        if not sections["education"]:

            recommendations.append(
                "Add your education details."
            )

        if not sections["summary"]:

            recommendations.append(
                "Add a professional summary."
            )

        if "python" not in skills_list:

            recommendations.append(
                "Consider adding Python if relevant."
            )

        if "git" not in skills_list:

            recommendations.append(
                "Add Git/GitHub experience."
            )

        if "sql" not in skills_list:

            recommendations.append(
                "Add SQL/database skills if relevant."
            )

        if not recommendations:

            recommendations.append(
                "Your resume structure looks strong. "
                "Keep improving your projects, "
                "skills and experience."
            )

        for recommendation in recommendations:

            print(
                f"• {recommendation}"
            )

        print(
            "\n========================================"
        )

    # ============================================================
    # ATS SCORE
    # ============================================================

    def ats_score(
        self
    ):

        print(
            "\n========== ATS SCORE =========="
        )

        print(
            "\nChoose Resume Input:"
        )

        print(
            "1. Parse Resume File"
        )

        print(
            "2. Enter Skills Manually"
        )

        choice = input(
            "\nEnter your choice: "
        ).strip()

        # ========================================================
        # RESUME FILE
        # ========================================================

        if choice == "1":

            file_path = input(
                "\nEnter resume file path: "
            ).strip()

            file_path = file_path.strip(
                '"'
            )

            resume_text = (
                self.parse_resume_file(
                    file_path
                )
            )

            if not resume_text:

                return

            print(
                "\n✅ Resume parsed successfully!"
            )

        # ========================================================
        # MANUAL SKILLS
        # ========================================================

        elif choice == "2":

            resume_text = input(
                "\nEnter your skills "
                "(comma separated): "
            )

        else:

            print(
                "\n❌ Invalid choice."
            )

            return

        # ========================================================
        # SKILLS
        # ========================================================

        skills = (
            self.extract_skills_from_resume(
                resume_text
            )
        )

        # ========================================================
        # SECTIONS
        # ========================================================

        sections = (
            self.detect_resume_sections(
                resume_text
            )
        )

        score = 0

        # ========================================================
        # SKILL SCORE
        # ========================================================

        if "python" in skills:

            score += 15

        if "git" in skills:

            score += 10

        if "sql" in skills:

            score += 10

        if (
            "machine learning" in skills
            or
            "artificial intelligence" in skills
            or
            "ai" in skills
        ):

            score += 10

        if (
            "pandas" in skills
            or
            "numpy" in skills
        ):

            score += 5

        # ========================================================
        # SECTION SCORE
        # ========================================================

        if sections["education"]:

            score += 10

        if sections["projects"]:

            score += 15

        if sections["experience"]:

            score += 10

        if sections["certifications"]:

            score += 5

        if sections["summary"]:

            score += 5

        # ========================================================
        # CONTACT SCORE
        # ========================================================

        email = self.extract_email(
            resume_text
        )

        phone = self.extract_phone(
            resume_text
        )

        if email != "Not detected":

            score += 3

        if phone != "Not detected":

            score += 2

        score = min(
            score,
            100
        )

        # ========================================================
        # RESULT
        # ========================================================

        print(
            "\n========== ATS RESULT =========="
        )

        print(
            f"\n🎯 ATS Score: {score}/100"
        )

        # Progress bar

        bar_length = 30

        filled = int(
            bar_length
            *
            score
            /
            100
        )

        progress_bar = (

            "█" * filled
            +
            "░" * (
                bar_length
                -
                filled
            )
        )

        print(
            f"[{progress_bar}]"
        )

        if score >= 80:

            print(
                "\n🔥 Excellent ATS Compatibility"
            )

        elif score >= 65:

            print(
                "\n✅ Good ATS Compatibility"
            )

        elif score >= 50:

            print(
                "\n👍 Average ATS Compatibility"
            )

        else:

            print(
                "\n⚠ Low ATS Compatibility"
            )

        # ========================================================
        # ATS RECOMMENDATIONS
        # ========================================================

        print(
            "\n💡 ATS Recommendations"
        )

        ats_recommendations = []

        if "python" not in skills:

            ats_recommendations.append(
                "Add Python if it is relevant to your target role."
            )

        if "git" not in skills:

            ats_recommendations.append(
                "Add Git/GitHub experience."
            )

        if "sql" not in skills:

            ats_recommendations.append(
                "Add SQL/database skills if relevant."
            )

        if not sections["projects"]:

            ats_recommendations.append(
                "Add projects with measurable results."
            )

        if not sections["experience"]:

            ats_recommendations.append(
                "Add internship/work experience."
            )

        if not sections["education"]:

            ats_recommendations.append(
                "Add education details."
            )

        if not sections["certifications"]:

            ats_recommendations.append(
                "Add relevant certifications."
            )

        if not ats_recommendations:

            ats_recommendations.append(
                "Your resume is well structured for ATS."
            )

        for recommendation in ats_recommendations:

            print(
                f"• {recommendation}"
            )

        print(
            "\n========================================"
        )

    # ============================================================
    # INTERVIEW PREPARATION
    # ============================================================

    def interview_preparation(
        self
    ):

        print(
            "\n========== Interview Preparation =========="
        )

        questions = [

            "What is Python?",

            "What is a List in Python?",

            "What is a Function?"
        ]

        answers = [

            "Python is a high-level, "
            "interpreted programming language.",

            "A List is a collection used "
            "to store multiple items.",

            "A Function is a reusable block "
            "of code that performs a specific task."
        ]

        score = 0

        for i in range(
            len(questions)
        ):

            print(
                f"\nQuestion {i + 1}"
            )

            print(
                questions[i]
            )

            user_answer = input(
                "Your Answer: "
            )

            if len(
                user_answer.strip()
            ) > 10:

                score += 1

            print(
                "\nSample Answer:"
            )

            print(
                answers[i]
            )

        print(
            f"\n🏆 Your Practice Score: "
            f"{score}/{len(questions)}"
        )

        print(
            "🎉 Interview Practice Completed!"
        )

    # ============================================================
    # CAREER HISTORY
    # ============================================================

    def career_history(
        self
    ):

        print(
            "\n========== Career History =========="
        )

        try:

            # Use the same directory as
            # the main application.

            base_dir = os.path.dirname(
                os.path.dirname(
                    os.path.abspath(
                        __file__
                    )
                )
            )

            data_file = os.path.join(
                base_dir,
                "career_data.json"
            )

            careers = load_dataset(data_file, [])

            found = False

            for career in careers:

                # Show only current user's records

                if (
                    career.get("user")
                    !=
                    self.current_user
                ):

                    continue

                print(
                    "\n------------------------------"
                )

                print(
                    f"Name       : "
                    f"{career.get('name', '')}"
                )

                print(
                    f"Career     : "
                    f"{career.get('career', '')}"
                )

                print(
                    f"Progress   : "
                    f"{career.get('progress', 0)}%"
                )

                if "timestamp" in career:

                    print(
                        f"Created On : "
                        f"{career['timestamp']}"
                    )

                if career.get(
                    "favorite",
                    False
                ):

                    print(
                        "⭐ Favorite"
                    )

                print(
                    "------------------------------"
                )

                found = True

            if not found:

                print(
                    "No career history found "
                    "for the current user."
                )

        except FileNotFoundError:

            print(
                "No career history found."
            )

        except Exception as e:

            print(
                f"❌ Error: {e}"
            )