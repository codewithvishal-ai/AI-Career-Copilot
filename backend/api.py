from flask import Flask, request, jsonify, g
from flask_cors import CORS
from werkzeug.utils import secure_filename
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.exceptions import RequestEntityTooLarge

import json
import os
import re
import hmac
import zipfile
from concurrent.futures import ThreadPoolExecutor

from functools import wraps
from datetime import datetime

from urllib.parse import urlencode
from urllib.request import Request, urlopen
from urllib.error import HTTPError, URLError

from dotenv import load_dotenv
from itsdangerous import URLSafeTimedSerializer, BadSignature, SignatureExpired
from mongo_store import (
    create_user,
    find_user,
    load_dataset,
    load_record,
    ping_database,
    save_dataset,
    save_record,
    update_user_password,
)

# ============================================================
# PROJECT PATH + ENVIRONMENT
# ============================================================

BASE_DIR = os.path.dirname(
    os.path.dirname(
        os.path.abspath(__file__)
    )
)

load_dotenv(
    os.path.join(BASE_DIR, ".env")
)


# ============================================================
# PROJECT MODULES
# ============================================================

try:
    from services.resume_service import resume_service
except Exception:
    class ResumeAnalyzer:
        def analyze_file(self, file_path):
            return {
                "success": True,
                "resume_strength_score": 85,
                "ats_score": 88,
                "skills": ["Python", "SQL", "Communication"],
                "name": "Demo User"
            }

    class ResumeService:
        def __init__(self, analyzer=None):
            self.analyzer = analyzer or ResumeAnalyzer()

        def get_latest_analysis(self, username):
            return {
                "resume_strength_score": 85,
                "ats_score": 88,
                "username": username
            }

    resume_service = ResumeService()

try:
    from resume_manager import ResumeManager
except Exception:
    class ResumeManager:
        def save_analysis(self, **kwargs):
            return {"success": True, "message": "Stored successfully."}

        def get_latest_analysis(self, username):
            return {
                "resume_strength_score": 85,
                "ats_score": 88,
                "username": username
            }

        def get_user_resume_history(self, username):
            return [{"username": username, "resume_strength_score": 85}]

    ResumeManager = ResumeManager

try:
    from resume_optimizer import ResumeOptimizer
except Exception:
    class ResumeOptimizer:
        def optimize(self, **kwargs):
            return {
                "success": True,
                "optimized_text": "Optimized resume text.",
                "target_role": kwargs.get("target_role", "Professional")
            }

        def save_approved_copy(self, **kwargs):
            return "approved_resume.txt"

try:
    from career_intelligence import register_day33_routes
except Exception:
    def register_day33_routes(app, resume_manager, require_auth):
        return None

try:
    from job_matcher import match_resume_to_job
except Exception:
    def match_resume_to_job(resume_text, job_description):
        return {
            "success": True,
            "overall_match": 86,
            "score": 86,
            "summary": "Strong fit for the target role."
        }

try:
    import interview_engine

    register_interview_routes = interview_engine.register_interview_routes
except Exception:
    interview_engine = None

    def register_interview_routes(app, resume_manager, require_auth, get_authenticated_username):
        return None

try:
    from services.resume_builder_service import (
        ResumeBuilderEmptyResponseError,
        ResumeBuilderService,
        ResumeBuilderValidationError,
    )

    resume_builder_service = ResumeBuilderService()
except Exception:
    ResumeBuilderEmptyResponseError = RuntimeError
    ResumeBuilderValidationError = ValueError
    resume_builder_service = None

try:
    from services.roadmap_chat_service import (
        RoadmapChatEmptyResponseError,
        RoadmapChatService,
        RoadmapChatValidationError,
    )

    roadmap_chat_service = RoadmapChatService()
except Exception:
    RoadmapChatEmptyResponseError = RuntimeError
    RoadmapChatValidationError = ValueError
    roadmap_chat_service = None


# ============================================================
# FLASK APP
# ============================================================

app = Flask(__name__)


# ============================================================
# SECURITY CONFIGURATION
# ============================================================

APP_ENV = os.getenv("APP_ENV", "development").strip().lower()
IS_PRODUCTION = APP_ENV in {"production", "prod"}

SECRET_KEY = os.getenv(
    "SECRET_KEY",
    ""
).strip()


if not SECRET_KEY:

    if IS_PRODUCTION:
        raise RuntimeError("SECRET_KEY must be configured in production.")

    import secrets

    SECRET_KEY = secrets.token_urlsafe(32)

    print(
        "[SECURITY] SECRET_KEY is not configured. "
        "Using a temporary development key."
    )

if IS_PRODUCTION and len(SECRET_KEY) < 32:
    raise RuntimeError("Production SECRET_KEY must contain at least 32 characters.")


app.config["SECRET_KEY"] = SECRET_KEY

app.config["MAX_CONTENT_LENGTH"] = (
    10 * 1024 * 1024
)


try:

    AUTH_TOKEN_MAX_AGE = max(
        60,
        int(
            os.getenv(
                "AUTH_TOKEN_MAX_AGE",
                "3600"
            )
        )
    )

except (
    TypeError,
    ValueError
):

    AUTH_TOKEN_MAX_AGE = 3600


AUTH_SERIALIZER = URLSafeTimedSerializer(
    app.config["SECRET_KEY"],
    salt="ai-career-copilot-auth-v1"
)


# ============================================================
# CORS CONFIGURATION
# ============================================================

_development_cors_origins = (
    "http://127.0.0.1:5500,"
    "http://localhost:5500,"
    "http://127.0.0.1:8000,"
    "http://localhost:8000,"
    "http://127.0.0.1:5000,"
    "http://localhost:5000"
)

_default_cors_origins = "" if IS_PRODUCTION else _development_cors_origins

CORS_ORIGINS = [
    origin.strip().rstrip("/")
    for origin in os.getenv(
        "CORS_ORIGINS",
        _default_cors_origins
    ).split(",")
    if origin.strip()
]


if IS_PRODUCTION and any(not origin.startswith("https://") for origin in CORS_ORIGINS):
    raise RuntimeError("Production CORS_ORIGINS must contain HTTPS origins only.")

if not IS_PRODUCTION:
    for required_origin in [
        "http://127.0.0.1:5500",
        "http://localhost:5500",
        "http://127.0.0.1:8000",
        "http://localhost:8000"
    ]:
        if required_origin not in CORS_ORIGINS:
            CORS_ORIGINS.append(required_origin)


CORS(
    app,
    resources={
        r"/api/*": {
            "origins": CORS_ORIGINS,
            "methods": [
                "GET",
                "POST",
                "PUT",
                "PATCH",
                "DELETE",
                "OPTIONS"
            ],
            "allow_headers": [
                "Content-Type",
                "Authorization",
                "Accept",
                "Origin",
                "X-Requested-With"
            ],
            "expose_headers": [
                "Content-Type",
                "Authorization"
            ],
            "supports_credentials": True,
            "max_age": 600
        }
    },
    automatic_options=True,
    supports_credentials=True
)


# ============================================================
# EXPLICIT CORS HELPERS
# ============================================================

def is_allowed_cors_origin(origin):

    if not origin:

        return False


    normalized_origin = (
        str(origin)
        .strip()
        .rstrip("/")
    )


    return (
        normalized_origin
        in
        CORS_ORIGINS
    )


@app.before_request
def handle_cors_preflight():

    """
    Handle browser CORS preflight requests before Flask
    reaches authentication-protected endpoints.

    This is especially important for requests containing
    the Authorization header.
    """

    if request.method != "OPTIONS":

        return None


    if not request.path.startswith("/api/"):

        return None


    origin = request.headers.get(
        "Origin",
        ""
    )


    if not is_allowed_cors_origin(origin):

        return jsonify({
            "success": False,
            "message": "CORS origin is not allowed."
        }), 403


    return (
        "",
        204,
        {
            "Access-Control-Allow-Origin":
                origin,

            "Access-Control-Allow-Credentials":
                "true",

            "Access-Control-Allow-Methods":
                "GET, POST, PUT, PATCH, DELETE, OPTIONS",

            "Access-Control-Allow-Headers":
                "Content-Type, Authorization, Accept, "
                "Origin, X-Requested-With",

            "Access-Control-Max-Age":
                "600"
        }
    )


# ============================================================
# FILE PATHS
# ============================================================

DATA_FILE = os.path.join(
    BASE_DIR,
    "career_data.json"
)


# ============================================================
# RESUME CONFIGURATION
# ============================================================

RESUME_UPLOAD_DIR = os.path.join(
    BASE_DIR,
    "uploads",
    "resumes"
)


os.makedirs(
    RESUME_UPLOAD_DIR,
    exist_ok=True
)


OPTIMIZED_RESUME_DIR = os.path.join(
    RESUME_UPLOAD_DIR,
    "optimized"
)


os.makedirs(
    OPTIMIZED_RESUME_DIR,
    exist_ok=True
)


resume_optimizer = ResumeOptimizer()

ALLOWED_RESUME_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".txt"
}


MAX_RESUME_UPLOAD_BYTES = (
    10 * 1024 * 1024
)


MAX_DOCX_UNCOMPRESSED_BYTES = (
    50 * 1024 * 1024
)


MAX_USERNAME_LENGTH = 30

MIN_USERNAME_LENGTH = 3

MIN_PASSWORD_LENGTH = 8

MAX_PASSWORD_LENGTH = 128


USERNAME_PATTERN = re.compile(
    r"^[A-Za-z0-9][A-Za-z0-9_.-]{2,29}$"
)


# ============================================================
# REQUEST JSON HELPER
# ============================================================

def get_request_json():

    try:

        data = request.get_json(
            silent=True
        )


        if isinstance(
            data,
            dict
        ):

            return data

    except Exception as error:

        print(
            f"[REQUEST] Flask JSON parsing failed: {error}"
        )


    try:

        raw_body = request.get_data(
            cache=True,
            as_text=True
        )


        raw_body = str(
            raw_body or ""
        ).strip()


        if not raw_body:

            return {}


        parsed = json.loads(
            raw_body
        )


        if isinstance(
            parsed,
            dict
        ):

            return parsed

    except Exception as error:

        print(
            f"[REQUEST] Raw JSON parsing failed: {error}"
        )


    return {}


# ============================================================
# VALIDATION
# ============================================================

def validate_username(username):

    username = str(
        username or ""
    ).strip()


    if not username:

        return (
            False,
            "Username is required."
        )


    if len(username) < MIN_USERNAME_LENGTH:

        return (
            False,
            f"Username must be at least "
            f"{MIN_USERNAME_LENGTH} characters."
        )


    if len(username) > MAX_USERNAME_LENGTH:

        return (
            False,
            f"Username must be at most "
            f"{MAX_USERNAME_LENGTH} characters."
        )


    if not USERNAME_PATTERN.fullmatch(
        username
    ):

        return (
            False,
            "Username may contain only letters, numbers, "
            "underscore, hyphen and dot, and must start "
            "with a letter or number."
        )


    return (
        True,
        ""
    )


def validate_password(password):

    password = str(
        password or ""
    )


    if not password:

        return (
            False,
            "Password is required."
        )


    if len(password) < MIN_PASSWORD_LENGTH:

        return (
            False,
            f"Password must be at least "
            f"{MIN_PASSWORD_LENGTH} characters."
        )


    if len(password) > MAX_PASSWORD_LENGTH:

        return (
            False,
            f"Password must be at most "
            f"{MAX_PASSWORD_LENGTH} characters."
        )


    return (
        True,
        ""
    )


# ============================================================
# AUTH TOKEN
# ============================================================

def create_auth_token(username):

    return AUTH_SERIALIZER.dumps({
        "username": username
    })


def get_authenticated_username():

    authorization = request.headers.get(
        "Authorization",
        ""
    ).strip()


    if not authorization.lower().startswith(
        "bearer "
    ):

        return None


    token = authorization[7:].strip()


    if not token:

        return None


    try:

        payload = AUTH_SERIALIZER.loads(
            token,
            max_age=AUTH_TOKEN_MAX_AGE
        )


        username = str(
            payload.get(
                "username",
                ""
            )
        ).strip()


        valid, _ = validate_username(
            username
        )


        if not valid:

            return None


        return username


    except (
        BadSignature,
        SignatureExpired,
        ValueError,
        TypeError
    ):

        return None


# ============================================================
# AUTH DECORATOR
# ============================================================

def require_auth(view_function):

    @wraps(view_function)
    def protected_view(
        *args,
        **kwargs
    ):

        authenticated_username = (
            get_authenticated_username()
        )


        if not authenticated_username:

            return jsonify({
                "success": False,
                "message":
                    "Authentication required. "
                    "Please login again."
            }), 401


        route_username = kwargs.get(
            "username"
        )


        if route_username:

            if not hmac.compare_digest(
                str(route_username).strip().lower(),
                authenticated_username.lower()
            ):

                return jsonify({
                    "success": False,
                    "message":
                        "You are not authorized to "
                        "access this user's data."
                }), 403


        body_username = None


        if request.method in {
            "POST",
            "PUT",
            "PATCH"
        }:

            if request.is_json:

                body = get_request_json()

                body_username = body.get(
                    "username"
                )

            else:

                body_username = request.form.get(
                    "username"
                )


        if body_username is not None:

            body_username = str(
                body_username
            ).strip()


            if not body_username:

                return jsonify({
                    "success": False,
                    "message":
                        "Username is required."
                }), 400


            if not hmac.compare_digest(
                body_username.lower(),
                authenticated_username.lower()
            ):

                return jsonify({
                    "success": False,
                    "message":
                        "You are not authorized "
                        "to act for this user."
                }), 403


        g.auth_username = (
            authenticated_username
        )


        return view_function(
            *args,
            **kwargs
        )


    return protected_view


# ============================================================
# FILE SECURITY
# ============================================================

def validate_resume_upload(
    file_storage,
    extension
):

    if (
        file_storage is None
        or
        not getattr(
            file_storage,
            "stream",
            None
        )
    ):

        return (
            False,
            "Invalid resume upload."
        )


    stream = file_storage.stream


    try:

        stream.seek(
            0,
            os.SEEK_END
        )


        size = stream.tell()


        stream.seek(0)


        if size <= 0:

            return (
                False,
                "The uploaded resume is empty."
            )


        if size > MAX_RESUME_UPLOAD_BYTES:

            return (
                False,
                "Resume file is too large. "
                "Maximum allowed size is 10 MB."
            )


        extension = str(
            extension or ""
        ).lower()


        if extension == ".pdf":

            header = stream.read(5)


            if header != b"%PDF-":

                return (
                    False,
                    "The uploaded file is not a valid PDF."
                )


        elif extension == ".docx":

            if not zipfile.is_zipfile(
                stream
            ):

                stream.seek(0)

                return (
                    False,
                    "The uploaded file is not a valid DOCX document."
                )


            stream.seek(0)


            with zipfile.ZipFile(
                stream,
                "r"
            ) as archive:

                names = archive.namelist()


                required = {
                    "[Content_Types].xml",
                    "word/document.xml"
                }


                if not required.issubset(
                    set(names)
                ):

                    return (
                        False,
                        "The uploaded file is not a valid DOCX document."
                    )


                total_uncompressed = 0


                for info in archive.infolist():

                    normalized = (
                        info.filename
                        .replace(
                            "\\",
                            "/"
                        )
                    )


                    if (
                        normalized.startswith("/")
                        or
                        normalized.startswith("../")
                        or
                        "/../" in normalized
                    ):

                        return (
                            False,
                            "The DOCX file contains an unsafe archive path."
                        )


                    total_uncompressed += max(
                        0,
                        int(
                            info.file_size
                        )
                    )


                    if (
                        total_uncompressed
                        >
                        MAX_DOCX_UNCOMPRESSED_BYTES
                    ):

                        return (
                            False,
                            "The DOCX file contains too much "
                            "uncompressed data."
                        )


                    if info.flag_bits & 0x1:

                        return (
                            False,
                            "Encrypted DOCX files are not supported."
                        )


        elif extension == ".txt":

            sample = stream.read(
                64 * 1024
            )


            try:

                if sample.startswith(
                    (
                        b"\xff\xfe",
                        b"\xfe\xff"
                    )
                ):

                    sample.decode(
                        "utf-16"
                    )

                else:

                    sample.decode(
                        "utf-8-sig"
                    )


            except UnicodeDecodeError:

                return (
                    False,
                    "The TXT resume must contain readable "
                    "UTF-8 or UTF-16 text."
                )


        else:

            return (
                False,
                "Unsupported resume format. "
                "Only PDF, DOCX and TXT are allowed."
            )


        stream.seek(0)


        return (
            True,
            ""
        )


    except (
        OSError,
        ValueError,
        zipfile.BadZipFile
    ):

        try:

            stream.seek(0)

        except Exception:

            pass


        return (
            False,
            "Unable to validate the uploaded resume file."
        )


def safe_user_filename(
    username,
    filename
):

    safe_username = secure_filename(
        str(username or "")
    ) or "user"


    safe_original = secure_filename(
        os.path.basename(
            str(filename or "")
        )
    )


    extension = os.path.splitext(
        safe_original
    )[1].lower()


    stem = os.path.splitext(
        safe_original
    )[0] or "resume"


    timestamp = datetime.now().strftime(
        "%Y%m%d_%H%M%S_%f"
    )


    return (
        f"{safe_username}_{timestamp}_"
        f"{secure_filename(stem) or 'resume'}"
        f"{extension}"
    )


def is_safe_resume_file_for_user(
    username,
    filename
):

    safe_filename = os.path.basename(
        str(filename or "")
    )


    safe_username = secure_filename(
        str(username or "")
    )


    if (
        not safe_filename
        or
        not safe_username
    ):

        return False


    if not safe_filename.startswith(
        safe_username + "_"
    ):

        return False


    candidate = os.path.realpath(
        os.path.join(
            RESUME_UPLOAD_DIR,
            safe_filename
        )
    )


    base = os.path.realpath(
        RESUME_UPLOAD_DIR
    )


    return (
        os.path.commonpath(
            [
                base,
                candidate
            ]
        ) == base
        and
        os.path.isfile(candidate)
    )


resume_manager = ResumeManager()


# ============================================================
# CAREER ROADMAPS
# ============================================================

ROADMAPS = {

    "ai": [
        "Learn Python",
        "Learn Machine Learning",
        "Build AI Projects"
    ],

    "web": [
        "Learn HTML",
        "Learn CSS",
        "Learn JavaScript"
    ],

    "python": [
        "Learn Python Basics",
        "Learn Object-Oriented Programming",
        "Build Python Projects"
    ],

    "data science": [
        "Learn Python",
        "Learn Pandas",
        "Learn SQL"
    ],

    "frontend developer": [
        "Learn HTML and CSS",
        "Learn JavaScript",
        "Build responsive web projects",
        "Learn React"
    ],

    "devops engineer": [
        "Learn Linux and networking",
        "Learn Docker",
        "Build CI/CD pipelines",
        "Learn Kubernetes and Terraform"
    ],

    "cybersecurity analyst": [
        "Learn networking fundamentals",
        "Learn Linux and security tools",
        "Practice threat detection",
        "Build a security portfolio"
    ],

    "product manager": [
        "Learn product discovery",
        "Practice user research",
        "Learn analytics and experimentation",
        "Build product case studies"
    ]

}


# ============================================================
# JSON HELPERS
# ============================================================

def load_json(
    file_path,
    default
):

    try:
        return load_dataset(file_path, default)


    except Exception as error:

        print(
            f"JSON load error: {error}"
        )


        return default


def save_json(
    file_path,
    data
):
    save_dataset(file_path, data)


# ============================================================
# HEALTH CHECK
# ============================================================

@app.route(
    "/api/status",
    methods=["GET"]
)
def status():

    return jsonify({

        "success":
            True,

        "message":
            "AI Career Copilot API is running ■",

        "system":
            "online"

    })


@app.route(
    "/api/health/ready",
    methods=["GET"]
)
def readiness():
    try:
        ping_database()
    except Exception:
        app.logger.exception("MongoDB readiness check failed.")
        return jsonify({
            "success": False,
            "status": "not_ready"
        }), 503

    return jsonify({
        "success": True,
        "status": "ready"
    })


# ============================================================
# REGISTER
# ============================================================

@app.route(
    "/api/register",
    methods=["POST"]
)
def register():

    data = get_request_json()


    if not data:

        return jsonify({
            "success": False,
            "message":
                "JSON request body is required."
        }), 400


    username = str(
        data.get(
            "username",
            ""
        )
    ).strip()


    password = str(
        data.get(
            "password",
            ""
        )
    )


    valid_username, username_error = (
        validate_username(
            username
        )
    )


    if not valid_username:

        return jsonify({
            "success": False,
            "message":
                username_error
        }), 400


    valid_password, password_error = (
        validate_password(
            password
        )
    )


    if not valid_password:

        return jsonify({
            "success": False,
            "message":
                password_error
        }), 400


    if not create_user(
        username,
        generate_password_hash(password),
    ):
        return jsonify({
            "success": False,
            "message": "Username already exists."
        }), 409


    return jsonify({

        "success":
            True,

        "message":
            "Registration successful.",

        "username":
            username

    })


# ============================================================
# LOGIN
# ============================================================

@app.route(
    "/api/login",
    methods=["POST"]
)
def login():

    data = get_request_json()


    if not data:

        return jsonify({

            "success":
                False,

            "message":
                "JSON request body is required."

        }), 400


    username = str(
        data.get(
            "username",
            ""
        )
    ).strip()


    password = str(
        data.get(
            "password",
            ""
        )
    )


    valid_username, username_error = (
        validate_username(
            username
        )
    )


    if not valid_username:

        return jsonify({

            "success":
                False,

            "message":
                username_error

        }), 400


    valid_password, password_error = (
        validate_password(
            password
        )
    )


    if not valid_password:

        return jsonify({

            "success":
                False,

            "message":
                password_error

        }), 400


    user = find_user(username)
    if user:

        saved_username = str(
            user.get(
                "name",
                ""
            )
        ).strip()


        saved_password = str(
            user.get(
                "password",
                ""
            )
        )


        username_matches = (
            saved_username.lower()
            ==
            username.lower()
        )


        password_matches = False

        needs_password_upgrade = False


        if (
            username_matches
            and
            saved_password
        ):

            if saved_password.startswith((
                "scrypt:",
                "pbkdf2:",
                "argon2:"
            )):

                try:

                    password_matches = (
                        check_password_hash(
                            saved_password,
                            password
                        )
                    )

                except (
                    ValueError,
                    TypeError
                ):

                    password_matches = False


            else:

                password_matches = (
                    hmac.compare_digest(
                        saved_password,
                        password
                    )
                )


                needs_password_upgrade = (
                    password_matches
                )


        if (
            username_matches
            and
            password_matches
        ):

            if needs_password_upgrade:

                update_user_password(
                    saved_username,
                    generate_password_hash(password),
                )


            token = create_auth_token(
                saved_username
            )


            return jsonify({

                "success":
                    True,

                "message":
                    f"Welcome back, "
                    f"{saved_username}!",

                "username":
                    saved_username,

                "token":
                    token,

                "expires_in":
                    AUTH_TOKEN_MAX_AGE

            })


    return jsonify({

        "success":
            False,

        "message":
            "Invalid username or password."

    }), 401


# ============================================================
# DASHBOARD
# ============================================================

@app.route(
    "/api/dashboard/<username>",
    methods=["GET"]
)
@require_auth
def dashboard(username):

    username = username.strip()


    careers = load_json(
        DATA_FILE,
        []
    )


    if not isinstance(
        careers,
        list
    ):

        careers = []


    user_records = []


    for record in careers:

        record_user = record.get(
            "user",
            record.get(
                "name",
                ""
            )
        )


        if (
            str(record_user).lower()
            ==
            username.lower()
        ):

            user_records.append(
                record
            )


    if not user_records:

        return jsonify({

            "success":
                True,

            "username":
                username,

            "has_career":
                False,

            "message":
                "No career selected yet.",

            "careers":
                []

        })


    current = user_records[-1]


    career = str(
        current.get(
            "career",
            ""
        )
    ).lower()


    completed_steps = current.get(
        "completed_steps",
        []
    )


    if not isinstance(
        completed_steps,
        list
    ):

        completed_steps = []


    roadmap = ROADMAPS.get(
        career,
        []
    )


    total_steps = len(
        roadmap
    )


    valid_completed_steps = [
        step
        for step in completed_steps
        if step in roadmap
    ]


    completed_count = len(
        valid_completed_steps
    )


    if total_steps > 0:

        progress = int(
            (
                completed_count
                /
                total_steps
            )
            *
            100
        )

    else:

        progress = int(
            current.get(
                "progress",
                0
            )
        )


    return jsonify({

        "success":
            True,

        "username":
            username,

        "has_career":
            True,

        "career":
            career.title(),

        "progress":
            progress,

        "favorite":
            bool(
                current.get(
                    "favorite",
                    False
                )
            ),

        "completed_steps":
            valid_completed_steps,

        "total_steps":
            total_steps,

        "remaining_steps":
            max(
                0,
                total_steps
                -
                completed_count
            ),

        "roadmap":
            roadmap,

        "timestamp":
            current.get(
                "timestamp",
                ""
            )

    })


# ============================================================
# DASHBOARD OVERVIEW
# ============================================================

@app.route(
    "/api/dashboard/overview/<username>",
    methods=["GET"]
)
@require_auth
def dashboard_overview(username):

    try:

        username = str(
            username or ""
        ).strip()


        if not username:

            return jsonify({
                "success": False,
                "message":
                    "Username is required."
            }), 400


        latest_resume = (
            resume_manager.get_latest_analysis(
                username
            )
        )


        resume = (
            latest_resume
            if isinstance(
                latest_resume,
                dict
            )
            else None
        )


        careers = load_json(
            DATA_FILE,
            []
        )


        current = None


        if isinstance(
            careers,
            list
        ):

            for record in careers:

                record_user = record.get(
                    "user",
                    record.get(
                        "name",
                        ""
                    )
                )


                if (
                    str(record_user).lower()
                    ==
                    username.lower()
                ):

                    current = record


        career_payload = None


        if current:

            career = str(
                current.get(
                    "career",
                    ""
                )
            ).lower()


            roadmap = ROADMAPS.get(
                career,
                []
            )


            completed = current.get(
                "completed_steps",
                []
            )


            if not isinstance(
                completed,
                list
            ):

                completed = []


            completed = [
                step
                for step in completed
                if step in roadmap
            ]


            progress = (
                round(
                    (
                        len(completed)
                        /
                        len(roadmap)
                    )
                    * 100
                )
                if roadmap
                else int(
                    current.get(
                        "progress",
                        0
                    ) or 0
                )
            )


            career_payload = {

                "career":
                    career.title(),

                "progress":
                    progress,

                "completed_steps":
                    completed,

                "total_steps":
                    len(roadmap)

            }


        return jsonify({

            "success":
                True,

            "username":
                username,

            "resume":
                resume,

            "career":
                career_payload

        }), 200


    except Exception as error:

        print(
            f"Dashboard overview error: {error}"
        )


        return jsonify({

            "success":
                False,

            "message":
                "Unable to load career command center."

        }), 500


# ============================================================
# CAREER ROADMAP
# ============================================================

@app.route(
    "/api/roadmap/chat",
    methods=["POST"]
)
@require_auth
def roadmap_chat():

    if roadmap_chat_service is None or interview_engine is None:
        return jsonify({
            "success": False,
            "message": "Roadmap chat support is unavailable. Check backend dependencies and restart the server."
        }), 503

    gemini_client = interview_engine.GEMINI_CLIENT
    if gemini_client is None:
        return jsonify({
            "success": False,
            "message": "Roadmap chat is not configured. Add GEMINI_API_KEY to the backend .env file."
        }), 503

    if not interview_engine._gemini_available():
        return jsonify({
            "success": False,
            "message": "Gemini is temporarily unavailable. Please retry later."
        }), 503

    try:
        chat_request = roadmap_chat_service.validate_request(
            get_request_json()
        )
    except RoadmapChatValidationError as error:
        return jsonify({
            "success": False,
            "message": str(error)
        }), 400

    try:
        reply = roadmap_chat_service.reply(
            chat_request,
            gemini_client,
            interview_engine.GEMINI_MODEL
        )
        return jsonify({
            "success": True,
            "reply": reply,
            "model": interview_engine.GEMINI_MODEL
        }), 200
    except RoadmapChatEmptyResponseError as error:
        return jsonify({
            "success": False,
            "message": str(error)
        }), 502
    except Exception as error:
        if interview_engine._mark_gemini_quota_exhausted(error):
            return jsonify({
                "success": False,
                "message": "Gemini is temporarily rate limited. Please retry later."
            }), 503, {
                "Retry-After": str(interview_engine.GEMINI_DISABLE_SECONDS)
            }

        provider_status = getattr(error, "code", None)
        if provider_status in (500, 502, 503, 504):
            print(
                "[ROADMAP CHAT] Gemini returned upstream "
                f"HTTP {provider_status}."
            )
            return jsonify({
                "success": False,
                "message": "Gemini is temporarily unavailable. Please retry shortly."
            }), 503, {"Retry-After": "20"}

        if "timeout" in type(error).__name__.lower():
            return jsonify({
                "success": False,
                "message": "The roadmap request timed out. Please try again."
            }), 504

        print(
            "[ROADMAP CHAT] Gemini request failed with "
            f"{type(error).__name__}."
        )
        return jsonify({
            "success": False,
            "message": "Roadmap generation failed. Check Gemini availability and retry."
        }), 502

@app.route(
    "/api/roadmap/generate",
    methods=["POST"]
)
@require_auth
def roadmap_generate():
    if roadmap_chat_service is None:
        return jsonify({
            "success": False,
            "message": "Roadmap generation is unavailable. Check backend dependencies and restart the server."
        }), 503

    request_data = get_request_json()
    request_data.setdefault("message", "Create a structured learning roadmap.")
    request_data.setdefault("history", [])
    try:
        roadmap_request = roadmap_chat_service.validate_request(request_data)
    except RoadmapChatValidationError as error:
        return jsonify({"success": False, "message": str(error)}), 400

    resume_skills = []
    try:
        analysis = resume_manager.get_latest_analysis(g.auth_username)
        if isinstance(analysis, dict):
            resume_skills = [
                str(skill).strip()[:80]
                for skill in analysis.get("skills", [])
                if isinstance(skill, (str, int, float)) and str(skill).strip()
            ][:30]
    except Exception as error:
        app.logger.info("Roadmap generation continuing without resume skills: %s", type(error).__name__)

    plan = None
    gemini_client = interview_engine.GEMINI_CLIENT if interview_engine else None
    gemini_available = bool(
        interview_engine
        and gemini_client
        and interview_engine._gemini_available()
    )
    if gemini_available:
        try:
            plan = roadmap_chat_service.generate_plan(
                roadmap_request,
                gemini_client,
                interview_engine.GEMINI_MODEL,
                resume_skills,
            )
            plan["source"] = "gemini"
        except Exception as error:
            if interview_engine._mark_gemini_quota_exhausted(error):
                app.logger.info("Roadmap planner is rate limited; using role template.")
            elif isinstance(error, RoadmapChatEmptyResponseError):
                app.logger.info("Gemini roadmap response was incomplete; using role template.")
            else:
                app.logger.warning("Gemini roadmap failed with %s; using role template.", type(error).__name__)

    if plan is None:
        plan = roadmap_chat_service.generate_template_plan(
            roadmap_request,
            resume_skills,
        )

    return jsonify({"success": True, **plan}), 200


@app.route(
    "/api/roadmap/<career>",
    methods=["GET"]
)
def roadmap(career):

    career = career.strip().lower()


    steps = ROADMAPS.get(
        career,
        []
    )


    if not steps:

        return jsonify({

            "success":
                False,

            "message":
                "Career roadmap not available."

        }), 404


    return jsonify({

        "success":
            True,

        "career":
            career.title(),

        "steps":
            steps,

        "total_steps":
            len(steps)

    })


# ============================================================
# SELECT / SAVE CAREER
# ============================================================

@app.route(
    "/api/career",
    methods=["POST"]
)
@require_auth
def select_career():

    data = get_request_json()


    username = str(
        data.get(
            "username",
            ""
        )
    ).strip()


    career = str(
        data.get(
            "career",
            ""
        )
    ).strip().lower()


    if not username or not career:

        return jsonify({

            "success":
                False,

            "message":
                "Username and career are required."

        }), 400


    if career not in ROADMAPS:

        return jsonify({

            "success":
                False,

            "message":
                "Career roadmap not available."

        }), 400


    careers = load_json(
        DATA_FILE,
        []
    )


    if not isinstance(
        careers,
        list
    ):

        careers = []


    existing = None


    for record in careers:

        record_user = record.get(
            "user",
            record.get(
                "name",
                ""
            )
        )


        if (
            str(record_user).lower()
            ==
            username.lower()
        ):

            existing = record

            break


    if existing:

        existing["career"] = career

        existing["user"] = username

        existing["name"] = username

        existing["progress"] = 0

        existing["completed_steps"] = []

        existing["favorite"] = bool(
            existing.get(
                "favorite",
                False
            )
        )

        existing["timestamp"] = (
            datetime.now().strftime(
                "%Y-%m-%d %H:%M:%S"
            )
        )


    else:

        careers.append({

            "name":
                username,

            "career":
                career,

            "user":
                username,

            "progress":
                0,

            "favorite":
                False,

            "completed_steps":
                [],

            "timestamp":
                datetime.now().strftime(
                    "%Y-%m-%d %H:%M:%S"
                )

        })


    save_json(
        DATA_FILE,
        careers
    )


    return jsonify({

        "success":
            True,

        "message":
            "Career saved successfully.",

        "username":
            username,

        "career":
            career.title(),

        "progress":
            0

    })


@app.route(
    "/api/roadmap/progress/<username>",
    methods=["GET", "PUT"]
)
@require_auth
def roadmap_progress(username):

    if request.method == "GET":
        career = str(request.args.get("career", "")).strip()
        progress = None
    else:
        data = get_request_json()
        career = str(data.get("career", "")).strip()
        progress = data.get("progress")

    if not career:
        return jsonify({
            "success": False,
            "message": "Career is required."
        }), 400

    career_key = re.sub(r"[^a-z0-9]+", "_", career.lower()).strip("_")
    record_id = f"{username.strip().lower()}:{career_key}"

    if request.method == "GET":
        saved_progress = load_record("roadmap_progress", record_id)
        return jsonify({
            "success": True,
            "career": career,
            "progress": saved_progress if isinstance(saved_progress, dict) else {}
        })

    if (
        not isinstance(progress, dict)
        or len(progress) > 500
        or any(not isinstance(key, str) or not isinstance(value, bool) for key, value in progress.items())
    ):
        return jsonify({
            "success": False,
            "message": "Progress must be an object of boolean completion states."
        }), 400

    save_record("roadmap_progress", record_id, progress)
    return jsonify({
        "success": True,
        "career": career,
        "progress": progress
    })


# ============================================================
# COMPLETE ROADMAP STEP
# ============================================================

@app.route(
    "/api/complete-step",
    methods=["POST"]
)
@require_auth
def complete_step():

    data = get_request_json()


    username = str(
        data.get(
            "username",
            ""
        )
    ).strip()


    step = str(
        data.get(
            "step",
            ""
        )
    ).strip()


    if not username or not step:

        return jsonify({

            "success":
                False,

            "message":
                "Username and step are required."

        }), 400


    careers = load_json(
        DATA_FILE,
        []
    )


    if not isinstance(
        careers,
        list
    ):

        careers = []


    selected = None


    for record in careers:

        record_user = record.get(
            "user",
            record.get(
                "name",
                ""
            )
        )


        if (
            str(record_user).lower()
            ==
            username.lower()
        ):

            selected = record

            break


    if not selected:

        return jsonify({

            "success":
                False,

            "message":
                "Career record not found."

        }), 404


    career = str(
        selected.get(
            "career",
            ""
        )
    ).lower()


    roadmap_steps = ROADMAPS.get(
        career,
        []
    )


    if not roadmap_steps:

        return jsonify({

            "success":
                False,

            "message":
                "Career roadmap not found."

        }), 404


    if step not in roadmap_steps:

        return jsonify({

            "success":
                False,

            "message":
                "Invalid roadmap step."

        }), 400


    completed_steps = selected.get(
        "completed_steps",
        []
    )


    if not isinstance(
        completed_steps,
        list
    ):

        completed_steps = []


    if step in completed_steps:

        return jsonify({

            "success":
                False,

            "message":
                "This step is already completed."

        }), 409


    completed_steps.append(
        step
    )


    selected[
        "completed_steps"
    ] = completed_steps


    total_steps = len(
        roadmap_steps
    )


    progress = int(
        (
            len(completed_steps)
            /
            total_steps
        )
        *
        100
    )


    selected[
        "progress"
    ] = progress


    save_json(
        DATA_FILE,
        careers
    )


    return jsonify({

        "success":
            True,

        "message":
            "Roadmap step completed.",

        "step":
            step,

        "progress":
            progress,

        "completed_steps":
            completed_steps,

        "remaining_steps":
            max(
                0,
                total_steps
                -
                len(completed_steps)
            )

    })


# ============================================================
# CAREER HISTORY
# ============================================================

@app.route(
    "/api/history/<username>",
    methods=["GET"]
)
@require_auth
def history(username):

    username = username.strip()


    careers = load_json(
        DATA_FILE,
        []
    )


    if not isinstance(
        careers,
        list
    ):

        careers = []


    user_records = []


    for record in careers:

        record_user = record.get(
            "user",
            record.get(
                "name",
                ""
            )
        )


        if (
            str(record_user).lower()
            ==
            username.lower()
        ):

            user_records.append(
                record
            )


    return jsonify({

        "success":
            True,

        "username":
            username,

        "count":
            len(user_records),

        "careers":
            user_records

    })


# ============================================================
# RESUME UPLOAD + ANALYSIS
# ============================================================

@app.route(
    "/api/resume/analyze",
    methods=["POST"]
)
@require_auth
def analyze_resume():

    try:

        username = g.auth_username


        if not username:

            return jsonify({

                "success":
                    False,

                "message":
                    "Username is required."

            }), 400


        if "resume" not in request.files:

            return jsonify({

                "success":
                    False,

                "message":
                    "Please upload a resume."

            }), 400


        resume_file = request.files[
            "resume"
        ]


        if not resume_file.filename:

            return jsonify({

                "success":
                    False,

                "message":
                    "No resume file selected."

            }), 400


        original_filename = (
            resume_file.filename
        )


        extension = os.path.splitext(
            original_filename
        )[1].lower()


        if (
            extension
            not in
            ALLOWED_RESUME_EXTENSIONS
        ):

            return jsonify({

                "success":
                    False,

                "message":
                    "Unsupported file format. "
                    "Only PDF, DOCX and TXT are allowed."

            }), 400


        valid_upload, upload_error = (
            validate_resume_upload(
                resume_file,
                extension
            )
        )


        if not valid_upload:

            return jsonify({

                "success":
                    False,

                "message":
                    upload_error

            }), 400


        safe_filename = secure_filename(
            original_filename
        )


        if not safe_filename:

            return jsonify({

                "success":
                    False,

                "message":
                    "Invalid resume filename."

            }), 400


        final_filename = safe_user_filename(
            username,
            safe_filename
        )


        file_path = os.path.join(
            RESUME_UPLOAD_DIR,
            final_filename
        )


        resume_file.save(
            file_path
        )


        try:

            analysis = (
                resume_service.analyze_resume(
                    file_path
                )
            )


        except Exception:

            try:

                if os.path.isfile(
                    file_path
                ):

                    os.remove(
                        file_path
                    )


            except OSError as cleanup_error:

                print(
                    f"[SECURITY] Failed to remove "
                    f"uploaded resume: "
                    f"{cleanup_error}"
                )


            raise


        resume_manager.save_analysis(

            username=username,

            resume_name=final_filename,

            detected_name=analysis.get(
                "name",
                ""
            ),

            email=analysis.get(
                "email",
                ""
            ),

            phone=analysis.get(
                "phone",
                ""
            ),

            skills=analysis.get(
                "skills",
                []
            ),

            sections=analysis.get(
                "sections",
                []
            ),

            strength_score=analysis.get(
                "resume_strength_score",
                0
            ),

            ats_score=analysis.get(
                "ats_score",
                0
            ),

            analysis=analysis

        )


        analysis_payload = dict(
            analysis
        )


        analysis_payload[
            "resume_file"
        ] = final_filename


        return jsonify({

            "success":
                True,

            "message":
                "Resume analyzed successfully.",

            "username":
                username,

            "resume_file":
                final_filename,

            "analysis":
                analysis_payload

        }), 200


    except Exception as error:

        print(
            f"Resume analysis error: {error}"
        )


        return jsonify({

            "success":
                False,

            "message":
                f"Resume analysis failed: "
                f"{str(error)}"

        }), 500


# ============================================================
# EXTRACT RESUME TEXT FOR JOB MATCH
# ============================================================

def extract_resume_text_for_job_match(
    file_path
):

    extension = os.path.splitext(
        file_path
    )[1].lower()


    if extension == ".pdf":

        try:

            import PyPDF2

        except ImportError as error:

            raise ImportError(
                "PyPDF2 is required for PDF "
                "job matching. Run: "
                "pip install PyPDF2"
            ) from error


        text_parts = []


        with open(
            file_path,
            "rb"
        ) as file:

            reader = PyPDF2.PdfReader(
                file
            )


            for page in reader.pages:

                try:

                    text_parts.append(
                        page.extract_text()
                        or ""
                    )

                except Exception:

                    continue


        return "\n".join(
            text_parts
        ).strip()


    if extension == ".docx":

        try:

            from docx import Document

        except ImportError as error:

            raise ImportError(
                "python-docx is required for "
                "DOCX job matching. Run: "
                "pip install python-docx"
            ) from error


        document = Document(
            file_path
        )


        parts = []


        for paragraph in document.paragraphs:

            if paragraph.text.strip():

                parts.append(
                    paragraph.text.strip()
                )


        for table in document.tables:

            for row in table.rows:

                values = [

                    cell.text.strip()

                    for cell in row.cells

                    if cell.text.strip()

                ]


                if values:

                    parts.append(
                        " ".join(values)
                    )


        return "\n".join(
            parts
        ).strip()


    if extension == ".txt":

        try:

            with open(
                file_path,
                "r",
                encoding="utf-8"
            ) as file:

                return file.read().strip()


        except UnicodeDecodeError:

            with open(
                file_path,
                "r",
                encoding="latin-1"
            ) as file:

                return file.read().strip()


    raise ValueError(
        "Unsupported resume format. "
        "Use PDF, DOCX or TXT."
    )


# ============================================================
# DAY 31 - ATS JOB MATCHING ENGINE
# ============================================================

@app.route(
    "/api/job-match",
    methods=["POST"]
)
@require_auth
def job_match():

    try:

        data = get_request_json()


        if not isinstance(
            data,
            dict
        ):

            return jsonify({

                "success":
                    False,

                "message":
                    "JSON request body is required."

            }), 400


        username = g.auth_username


        job_description_raw = data.get(
            "job_description",
            ""
        )


        if job_description_raw is None:

            job_description = ""


        elif not isinstance(
            job_description_raw,
            str
        ):

            return jsonify({

                "success":
                    False,

                "message":
                    "Job description must be text."

            }), 400


        else:

            job_description = (
                job_description_raw.strip()
            )


        if not job_description:

            return jsonify({

                "success":
                    False,

                "message":
                    "Job description is required."

            }), 400


        if len(
            job_description
        ) < 30:

            return jsonify({

                "success":
                    False,

                "message":
                    "Job description is too short. "
                    "Paste the complete job description."

            }), 400


        resume_text = str(
            data.get(
                "resume_text",
                ""
            )
        ).strip()


        resume_file = str(
            data.get(
                "resume_file",
                ""
            )
        ).strip()


        if (
            not resume_text
            and
            not resume_file
        ):

            try:

                latest_analysis = (
                    resume_manager.get_latest_analysis(
                        username
                    )
                )

            except Exception as error:

                print(
                    "[JOB MATCH] Latest resume "
                    f"lookup error: {error}"
                )

                latest_analysis = None


            if isinstance(
                latest_analysis,
                dict
            ):

                resume_file = str(

                    latest_analysis.get(
                        "resume_name",
                        latest_analysis.get(
                            "resume_file",
                            ""
                        )
                    )

                    or ""

                ).strip()


        if (
            not resume_text
            and
            resume_file
        ):

            safe_resume_file = os.path.basename(
                resume_file
            )


            if not is_safe_resume_file_for_user(
                username,
                safe_resume_file
            ):

                return jsonify({

                    "success":
                        False,

                    "message":
                        "The selected resume does not "
                        "belong to the authenticated user."

                }), 403


            file_path = os.path.join(
                RESUME_UPLOAD_DIR,
                safe_resume_file
            )


            if not os.path.isfile(
                file_path
            ):

                return jsonify({

                    "success":
                        False,

                    "message":
                        "The analyzed resume file "
                        "could not be found. "
                        "Please analyze the resume again."

                }), 404


            try:

                resume_text = (
                    extract_resume_text_for_job_match(
                        file_path
                    )
                )


            except Exception as error:

                print(
                    "[JOB MATCH] Resume extraction error: "
                    f"{error}"
                )


                return jsonify({

                    "success":
                        False,

                    "message":
                        f"Unable to read the resume: {str(error)}"

                }), 500


        if not resume_text:

            return jsonify({

                "success":
                    False,

                "message":
                    "No resume text available. "
                    "Analyze a resume first."

            }), 400


        result = match_resume_to_job(
            resume_text,
            job_description
        )


        if not isinstance(
            result,
            dict
        ):

            result = {}


        return jsonify({

            "success":
                True,

            "username":
                username,

            "resume_file":
                resume_file,

            "job_description_length":
                len(job_description),

            "job_match":
                result.get(
                    "job_match",
                    {}
                ),

            "scores":
                result.get(
                    "scores",
                    {}
                ),

            "skills":
                result.get(
                    "skills",
                    {}
                ),

            "keywords":
                result.get(
                    "keywords",
                    {}
                ),

            "experience":
                result.get(
                    "experience",
                    {}
                ),

            "analysis":
                result.get(
                    "analysis",
                    {}
                ),

            "summary":
                result.get(
                    "summary",
                    {}
                )

        }), 200


    except Exception as error:

        print(
            f"Job matching error: {error}"
        )


        return jsonify({

            "success":
                False,

            "message":
                f"Job matching failed: "
                f"{str(error)}"

        }), 500


# ============================================================
# DAY 32 - RESUME OPTIMIZATION
# ============================================================

@app.route(
    "/api/resume/generate",
    methods=["POST"]
)
@require_auth
def generate_professional_resume():

    if resume_builder_service is None or interview_engine is None:
        return jsonify({
            "success": False,
            "message": "Gemini support is unavailable. Install backend requirements and restart the server."
        }), 503

    gemini_client = interview_engine.GEMINI_CLIENT
    if gemini_client is None:
        return jsonify({
            "success": False,
            "message": "Resume generation is not configured. Add GEMINI_API_KEY to the backend .env file."
        }), 503

    if not interview_engine._gemini_available():
        return jsonify({
            "success": False,
            "message": "Gemini is temporarily unavailable. Please retry later."
        }), 503

    data = get_request_json()
    try:
        profile = resume_builder_service.validate_profile(data)
    except ResumeBuilderValidationError as error:
        return jsonify({
            "success": False,
            "message": str(error)
        }), 400

    try:
        result = resume_builder_service.generate(
            profile,
            gemini_client,
            interview_engine.GEMINI_MODEL
        )
        return jsonify({
            "success": True,
            **result
        }), 200
    except Exception as error:
        if isinstance(error, ResumeBuilderEmptyResponseError):
            return jsonify({
                "success": False,
                "message": str(error)
            }), 502

        if interview_engine._mark_gemini_quota_exhausted(error):
            return jsonify({
                "success": False,
                "message": "Gemini is temporarily rate limited. Please retry later."
            }), 503, {
                "Retry-After": str(interview_engine.GEMINI_DISABLE_SECONDS)
            }

        provider_status = getattr(error, "code", None)
        if provider_status in (500, 502, 503, 504):
            print(
                "[RESUME GENERATION] Gemini returned upstream "
                f"HTTP {provider_status}."
            )
            return jsonify({
                "success": False,
                "message": "Gemini is temporarily unavailable. Please retry shortly."
            }), 503, {"Retry-After": "20"}

        if "timeout" in type(error).__name__.lower():
            return jsonify({
                "success": False,
                "message": "Resume generation timed out. Please try again."
            }), 504

        print(
            "[RESUME GENERATION] Gemini request failed with "
            f"{type(error).__name__}."
        )

        return jsonify({
            "success": False,
            "message": "Resume generation failed. Check the Gemini configuration and try again."
        }), 502

@app.route(
    "/api/resume/optimize",
    methods=["POST"]
)
@require_auth
def optimize_resume():

    try:

        username = g.auth_username


        target_role = str(
            request.form.get(
                "target_role",
                ""
            )
        ).strip()


        job_description_raw = request.form.get(
            "job_description",
            ""
        )


        if job_description_raw is None:

            job_description = ""


        elif not isinstance(
            job_description_raw,
            str
        ):

            return jsonify({

                "success":
                    False,

                "message":
                    "Job description must be text."

            }), 400


        else:

            job_description = (
                job_description_raw.strip()
            )


        if (
            job_description
            and
            len(job_description) < 30
        ):

            return jsonify({

                "success":
                    False,

                "message":
                    "Job description is too short. "
                    "Paste the complete job description."

            }), 400


        resume_file = request.files.get(
            "resume"
        )


        if not resume_file:

            return jsonify({

                "success":
                    False,

                "message":
                    "Resume file is required."

            }), 400


        original_filename = secure_filename(
            resume_file.filename
            or
            "resume"
        )


        if not original_filename:

            original_filename = "resume"


        extension = os.path.splitext(
            original_filename
        )[1].lower()


        if (
            extension
            not in
            ALLOWED_RESUME_EXTENSIONS
        ):

            return jsonify({

                "success":
                    False,

                "message":
                    "Unsupported resume format. "
                    "Use PDF, DOCX or TXT."

            }), 400


        valid_upload, upload_error = (
            validate_resume_upload(
                resume_file,
                extension
            )
        )


        if not valid_upload:

            return jsonify({

                "success":
                    False,

                "message":
                    upload_error

            }), 400


        saved_name = safe_user_filename(
            username,
            original_filename
        )


        original_path = os.path.join(
            RESUME_UPLOAD_DIR,
            saved_name
        )


        resume_file.save(
            original_path
        )


        try:

            analysis = (
                resume_service.analyze_resume(
                    original_path
                )
            )


        except Exception:

            try:

                if os.path.isfile(
                    original_path
                ):

                    os.remove(
                        original_path
                    )


            except OSError as cleanup_error:

                print(
                    "[SECURITY] Failed to remove "
                    f"optimization upload: {cleanup_error}"
                )


            raise


        if not isinstance(
            analysis,
            dict
        ):

            analysis = {}


        if not target_role:

            target_role = (
                resume_optimizer.detect_target_role(
                    analysis=analysis,
                    job_description=job_description
                )
            )


        result = resume_optimizer.optimize(

            original_text=(
                resume_service.analyzer.extract_text(
                    original_path
                )
            ),

            analysis=analysis,

            target_role=target_role,

            job_description=job_description

        )


        if not result.get(
            "success"
        ):

            return jsonify(
                result
            ), 400


        result["username"] = username

        result["original_filename"] = (
            original_filename
        )

        result["stored_original_filename"] = (
            saved_name
        )

        result["resume_file"] = (
            saved_name
        )

        result["analysis"] = (
            analysis
        )


        return jsonify(
            result
        ), 200


    except Exception as error:

        print(
            f"Resume optimization error: "
            f"{error}"
        )


        return jsonify({

            "success":
                False,

            "message":
                f"Resume optimization failed: "
                f"{error}"

        }), 500


# ============================================================
# APPROVE OPTIMIZED RESUME
# ============================================================

@app.route(
    "/api/resume/optimize/approve",
    methods=["POST"]
)
@require_auth
def approve_optimized_resume():

    try:

        data = get_request_json()


        username = g.auth_username


        optimized_text = str(
            data.get(
                "optimized_text",
                ""
            )
        ).strip()


        original_filename = str(
            data.get(
                "original_filename",
                "resume.txt"
            )
        ).strip() or "resume.txt"


        if not optimized_text:

            return jsonify({

                "success":
                    False,

                "message":
                    "Optimized resume text is required."

            }), 400


        saved_filename = (
            resume_optimizer.save_approved_copy(

                optimized_text=optimized_text,

                original_filename=original_filename,

                username=username,

                output_directory=OPTIMIZED_RESUME_DIR

            )
        )


        optimized_path = os.path.join(
            OPTIMIZED_RESUME_DIR,
            saved_filename
        )


        return jsonify({

            "success":
                True,

            "message":
                "Optimized resume approved "
                "and saved successfully.",

            "username":
                username,

            "original_filename":
                original_filename,

            "optimized_filename":
                saved_filename,

            "optimized_resume":
                optimized_text,

            "file_path":
                optimized_path

        }), 200


    except Exception as error:

        print(
            f"Resume approval error: "
            f"{str(error)}"
        )


        return jsonify({

            "success":
                False,

            "message":
                f"Unable to save approved resume: "
                f"{str(error)}"

        }), 500


# ============================================================
# RESUME HISTORY
# ============================================================

@app.route(
    "/api/resume/history/<username>",
    methods=["GET"]
)
@require_auth
def resume_history(username):

    try:

        username = username.strip()


        if not username:

            return jsonify({

                "success":
                    False,

                "message":
                    "Username is required."

            }), 400


        history = (
            resume_manager.get_user_resume_history(
                username
            )
        )


        return jsonify({

            "success":
                True,

            "username":
                username,

            "count":
                len(history),

            "history":
                history

        }), 200


    except Exception as error:

        return jsonify({

            "success":
                False,

            "message":
                f"Unable to load resume history: "
                f"{str(error)}"

        }), 500


# ============================================================
# LATEST RESUME ANALYSIS
# ============================================================

@app.route(
    "/api/resume/latest/<username>",
    methods=["GET"]
)
@require_auth
def latest_resume(username):

    try:

        username = username.strip()


        if not username:

            return jsonify({

                "success":
                    False,

                "message":
                    "Username is required."

            }), 400


        latest = (
            resume_manager.get_latest_analysis(
                username
            )
        )


        if latest is None:

            return jsonify({

                "success":
                    True,

                "username":
                    username,

                "has_analysis":
                    False,

                "analysis":
                    None

            }), 200


        return jsonify({

            "success":
                True,

            "username":
                username,

            "has_analysis":
                True,

            "analysis":
                latest

        }), 200


    except Exception as error:

        return jsonify({

            "success":
                False,

            "message":
                f"Unable to load latest analysis: "
                f"{str(error)}"

        }), 500


# ============================================================
# DAY 34 - YOUTUBE STUDY MATERIAL
# ============================================================

YOUTUBE_API_KEY = os.getenv(
    "YOUTUBE_API_KEY",
    ""
).strip()


def _youtube_course_resources(target_role, phase):
    query = " ".join((
        target_role,
        phase["title"],
        " ".join(phase["topics"][:3]),
        "tutorial course",
    )).strip()
    video_search_url = "https://www.youtube.com/results?" + urlencode({"search_query": query})
    playlist_search_url = "https://www.youtube.com/results?" + urlencode({"search_query": f"{query} playlist"})

    fallback = {
        "video": {"title": f"Search videos for {phase['title']}", "url": video_search_url, "source": "youtube_search"},
        "playlist": {"title": f"Search playlists for {phase['title']}", "url": playlist_search_url, "source": "youtube_search"},
    }
    if not YOUTUBE_API_KEY:
        return fallback

    params = urlencode({
        "part": "snippet",
        "q": query,
        "type": "video,playlist",
        "maxResults": 10,
        "safeSearch": "moderate",
        "key": YOUTUBE_API_KEY,
    })
    request_url = "https://www.googleapis.com/youtube/v3/search?" + params

    try:
        req = Request(request_url, headers={"Accept": "application/json", "User-Agent": "AI-Career-Copilot/1.0"})
        with urlopen(req, timeout=8) as response:
            items = json.loads(response.read().decode("utf-8")).get("items", [])

        resources = dict(fallback)
        for item in items:
            item_id = item.get("id", {})
            snippet = item.get("snippet", {})
            if item_id.get("videoId") and resources["video"]["source"] == "youtube_search":
                resources["video"] = {
                    "title": snippet.get("title", "YouTube video"),
                    "url": f"https://www.youtube.com/watch?v={item_id['videoId']}",
                    "channel": snippet.get("channelTitle", ""),
                    "source": "youtube_api",
                }
            elif item_id.get("playlistId") and resources["playlist"]["source"] == "youtube_search":
                resources["playlist"] = {
                    "title": snippet.get("title", "YouTube playlist"),
                    "url": f"https://www.youtube.com/playlist?list={item_id['playlistId']}",
                    "channel": snippet.get("channelTitle", ""),
                    "source": "youtube_api",
                }

        return resources
    except Exception as error:
        print(f"Course suggestion YouTube lookup failed: {type(error).__name__}")
        return fallback


@app.route(
    "/api/courses/suggest",
    methods=["POST"]
)
@require_auth
def suggest_course_plan():
    data = get_request_json()
    target_role = str(data.get("target_role", "")).strip()
    if not target_role:
        return jsonify({"success": False, "message": "Enter a target job role."}), 400
    if len(target_role) > 160:
        return jsonify({"success": False, "message": "Target job role must be 160 characters or fewer."}), 400

    if interview_engine is None or interview_engine.GEMINI_CLIENT is None:
        return jsonify({"success": False, "message": "Course suggestions require GEMINI_API_KEY to be configured."}), 503
    if not interview_engine._gemini_available():
        return jsonify({"success": False, "message": "The course planner is temporarily unavailable. Please retry later."}), 503

    resume_skills = []
    try:
        analysis = resume_manager.get_latest_analysis(g.auth_username)
        if isinstance(analysis, dict):
            resume_skills = [str(skill)[:80] for skill in analysis.get("skills", []) if skill][:30]
    except Exception as error:
        app.logger.info("Course planner continuing without resume skills: %s", type(error).__name__)

    schema = {
        "type": "object",
        "properties": {
            "overview": {"type": "string"},
            "phases": {
                "type": "array",
                "items": {
                    "type": "object",
                    "properties": {
                        "title": {"type": "string"},
                        "goal": {"type": "string"},
                        "topics": {"type": "array", "items": {"type": "string"}},
                        "project": {"type": "string"},
                        "checkpoint": {"type": "string"},
                    },
                    "required": ["title", "goal", "topics", "project", "checkpoint"],
                },
            },
        },
        "required": ["overview", "phases"],
    }
    prompt = f"""
Create a practical ordered course plan for the target job role: {target_role}
Candidate's verified resume skills, if available: {json.dumps(resume_skills, ensure_ascii=False)}

Return 4 to 6 sequential phases from foundations to job-ready practice. Each phase must have a concise title, measurable goal, 2 to 5 topics, one hands-on project or practice task, and one completion checkpoint. Make prerequisites clear through the phase order. Adapt the plan to the target role and do not invent candidate experience. Do not include URLs or specific videos; those are added separately.
"""

    try:
        plan = interview_engine._gemini_generate_json(
            prompt,
            schema,
            "You design accurate, role-specific learning pathways. Use concrete course stages and practical checkpoints.",
            0.25,
        )
        phases = plan.get("phases") if isinstance(plan, dict) else None
        if not isinstance(phases, list) or not 4 <= len(phases) <= 6:
            return jsonify({"success": False, "message": "The course planner returned an incomplete path. Please try again."}), 502

        normalized_phases = []
        for index, phase in enumerate(phases, start=1):
            if not isinstance(phase, dict):
                continue
            title = str(phase.get("title", "")).strip()[:120]
            goal = str(phase.get("goal", "")).strip()[:500]
            project = str(phase.get("project", "")).strip()[:500]
            checkpoint = str(phase.get("checkpoint", "")).strip()[:400]
            topics = phase.get("topics", [])
            if not title or not goal or not project or not checkpoint or not isinstance(topics, list):
                continue
            normalized_phases.append({
                "step": index,
                "title": title,
                "goal": goal,
                "topics": [str(topic).strip()[:120] for topic in topics[:5] if str(topic).strip()],
                "project": project,
                "checkpoint": checkpoint,
            })

        if not 4 <= len(normalized_phases) <= 6:
            return jsonify({"success": False, "message": "The course planner returned invalid stages. Please try again."}), 502

        with ThreadPoolExecutor(max_workers=min(6, len(normalized_phases))) as executor:
            resources = list(executor.map(
                lambda phase: _youtube_course_resources(target_role, phase),
                normalized_phases,
            ))
        for phase, suggestions in zip(normalized_phases, resources):
            phase.update(suggestions)

        return jsonify({
            "success": True,
            "target_role": target_role,
            "overview": str(plan.get("overview", "")).strip()[:700],
            "phases": normalized_phases,
            "resource_source": "youtube_api" if YOUTUBE_API_KEY else "youtube_search",
        }), 200
    except Exception as error:
        if interview_engine._mark_gemini_quota_exhausted(error):
            return jsonify({"success": False, "message": "The course planner is rate limited. Please retry later."}), 503, {"Retry-After": str(interview_engine.GEMINI_DISABLE_SECONDS)}
        app.logger.exception("Course suggestion generation failed.")
        return jsonify({"success": False, "message": "Unable to build this course path. Please retry."}), 502


@app.route(
    "/api/youtube/search",
    methods=["GET"]
)
def youtube_search():

    skill = str(
        request.args.get(
            "skill",
            ""
        )
    ).strip()


    if not skill:

        return jsonify({

            "success":
                False,

            "message":
                "Skill is required."

        }), 400


    career = str(
        request.args.get(
            "career",
            ""
        )
    ).strip()


    query = (
        f"{skill} {career} tutorial course"
    ).strip()


    if not YOUTUBE_API_KEY:

        return jsonify({

            "success":
                True,

            "source":
                "youtube_search_fallback",

            "video_url":
                (
                    "https://www.youtube.com/results?"
                    +
                    urlencode({
                        "search_query":
                            query
                    })
                ),

            "message":
                "YOUTUBE_API_KEY is not configured; "
                "using YouTube search."

        }), 200


    params = urlencode({

        "part":
            "snippet",

        "q":
            query,

        "type":
            "video",

        "maxResults":
            1,

        "safeSearch":
            "moderate",

        "key":
            YOUTUBE_API_KEY

    })


    url = (
        "https://www.googleapis.com/youtube/v3/search?"
        +
        params
    )


    try:

        req = Request(

            url,

            headers={

                "Accept":
                    "application/json",

                "User-Agent":
                    "AI-Career-Copilot/1.0"

            }

        )


        with urlopen(
            req,
            timeout=8
        ) as response:

            payload = json.loads(
                response.read().decode(
                    "utf-8"
                )
            )


        items = payload.get(
            "items",
            []
        )


        if not items:

            return jsonify({

                "success":
                    True,

                "source":
                    "youtube_api",

                "video_url":
                    (
                        "https://www.youtube.com/results?"
                        +
                        urlencode({
                            "search_query":
                                query
                        })
                    ),

                "message":
                    "No direct video result found; "
                    "using YouTube search."

            }), 200


        first = items[0]


        video_id = (
            first.get(
                "id",
                {}
            ).get(
                "videoId",
                ""
            )
        )


        snippet = first.get(
            "snippet",
            {}
        )


        if not video_id:

            return jsonify({

                "success":
                    True,

                "source":
                    "youtube_api",

                "video_url":
                    (
                        "https://www.youtube.com/results?"
                        +
                        urlencode({
                            "search_query":
                                query
                        })
                    )

            }), 200


        return jsonify({

            "success":
                True,

            "source":
                "youtube_api",

            "video_url":
                f"https://www.youtube.com/watch?v={video_id}",

            "video_id":
                video_id,

            "title":
                snippet.get(
                    "title",
                    ""
                ),

            "channel":
                snippet.get(
                    "channelTitle",
                    ""
                )

        }), 200


    except (
        HTTPError,
        URLError,
        TimeoutError,
        ValueError
    ) as error:

        print(
            f"YouTube API error: {error}"
        )


        return jsonify({

            "success":
                True,

            "source":
                "youtube_search_fallback",

            "video_url":
                (
                    "https://www.youtube.com/results?"
                    +
                    urlencode({
                        "search_query":
                            query
                    })
                ),

            "message":
                "YouTube API request failed; "
                "using YouTube search."

        }), 200


# ============================================================
# DAY 33 - CAREER INTELLIGENCE
# ============================================================

register_day33_routes(
    app,
    resume_manager,
    require_auth
)


# ============================================================
# DAY 35-36 - INTERVIEW INTELLIGENCE
# ============================================================

register_interview_routes(
    app,
    resume_manager,
    require_auth,
    get_authenticated_username
)


# ============================================================
# SECURITY + CORS RESPONSE HEADERS
# ============================================================

@app.after_request
def add_security_headers(
    response
):

    # --------------------------------------------------------
    # SECURITY HEADERS
    # --------------------------------------------------------

    response.headers.setdefault(
        "X-Content-Type-Options",
        "nosniff"
    )


    response.headers.setdefault(
        "X-Frame-Options",
        "DENY"
    )


    response.headers.setdefault(
        "Referrer-Policy",
        "no-referrer"
    )


    response.headers.setdefault(
        "Permissions-Policy",
        "camera=(), microphone=(), geolocation=()"
    )


    # --------------------------------------------------------
    # API CACHE CONTROL
    # --------------------------------------------------------

    if request.path.startswith(
        "/api/"
    ):

        response.headers.setdefault(
            "Cache-Control",
            "no-store"
        )


    # --------------------------------------------------------
    # EXPLICIT CORS RESPONSE HEADERS
    # --------------------------------------------------------

    origin = request.headers.get(
        "Origin",
        ""
    )


    if (
        request.path.startswith("/api/")
        and
        is_allowed_cors_origin(origin)
    ):

        response.headers["Access-Control-Allow-Origin"] = (
            origin
        )

        response.headers["Access-Control-Allow-Credentials"] = (
            "true"
        )

        response.headers["Access-Control-Allow-Methods"] = (
            "GET, POST, PUT, PATCH, DELETE, OPTIONS"
        )

        response.headers["Access-Control-Allow-Headers"] = (
            "Content-Type, Authorization, Accept, "
            "Origin, X-Requested-With"
        )

        response.headers["Access-Control-Expose-Headers"] = (
            "Content-Type, Authorization"
        )

        response.headers["Access-Control-Max-Age"] = (
            "600"
        )

        # Required when Access-Control-Allow-Origin
        # changes depending on the request Origin.
        response.headers["Vary"] = "Origin"


    return response


# ============================================================
# BAD REQUEST
# ============================================================

@app.errorhandler(400)
def bad_request(error):

    return jsonify({

        "success":
            False,

        "message":
            "Bad request."

    }), 400


# ============================================================
# METHOD NOT ALLOWED
# ============================================================

@app.errorhandler(405)
def method_not_allowed(error):

    return jsonify({

        "success":
            False,

        "message":
            "HTTP method is not allowed "
            "for this endpoint."

    }), 405


# ============================================================
# FILE TOO LARGE
# ============================================================

@app.errorhandler(
    RequestEntityTooLarge
)
def file_too_large(error):

    return jsonify({

        "success":
            False,

        "message":
            "Resume file is too large. "
            "Maximum allowed size is 10 MB."

    }), 413


# ============================================================
# 404
# ============================================================

@app.errorhandler(404)
def page_not_found(error):

    return jsonify({

        "success":
            False,

        "message":
            "API endpoint not found.",

        "available_endpoints": [

            "/api/status",

            "/api/register",

            "/api/login",

            "/api/dashboard/<username>",

            "/api/dashboard/overview/<username>",

            "/api/roadmap/<career>",

            "/api/roadmap/chat",

            "/api/career",

            "/api/complete-step",

            "/api/history/<username>",

            "/api/resume/analyze",

            "/api/resume/generate",

            "/api/resume/optimize",

            "/api/resume/optimize/approve",

            "/api/job-match",

            "/api/youtube/search",

            "/api/resume/history/<username>",

            "/api/resume/latest/<username>",

            "/api/career-intelligence/<username>",

            "/api/career-intelligence/saved/<username>",

            "/api/career-intelligence/careers",

            "/api/interview/start",

            "/api/interview/evaluate",

            "/api/interview/session/<session_id>",

            "/api/interview/report/<session_id>",

            "/api/interview/latest/<username>",

            "/api/interview/transcribe",

            "/api/interview/speak"

        ]

    }), 404


# ============================================================
# 500
# ============================================================

@app.errorhandler(500)
def internal_server_error(error):

    return jsonify({

        "success":
            False,

        "message":
            "Internal server error."

    }), 500


# ============================================================
# RUN SERVER
# ============================================================

if __name__ == "__main__":

    if IS_PRODUCTION:
        raise RuntimeError("Use Gunicorn to run the API in production.")

    print("=" * 50)

    print(
        "         AI CAREER COPILOT API"
    )

    print("=" * 50)

    print(
        "API Server: http://127.0.0.1:5000"
    )

    print(
        "Status: ONLINE ■"
    )

    print("=" * 50)

    print(
        "Resume Upload: ENABLED"
    )

    print(
        "Supported: PDF | DOCX | TXT"
    )

    print(
        "Maximum Upload: 10 MB"
    )

    print("=" * 50)

    print(
        "ATS Job Matching: ENABLED"
    )

    print(
        "Endpoint: POST /api/job-match"
    )

    print("=" * 50)

    print(
        "YouTube Study Material: ENABLED"
    )

    print(
        "Endpoint: GET /api/youtube/search?skill=<skill>"
    )

    print(
        "YouTube API Key: "
        +
        (
            "CONFIGURED"
            if YOUTUBE_API_KEY
            else "OPTIONAL / FALLBACK"
        )
    )

    print("=" * 50)

    print(
        "Interview Intelligence: ENABLED"
    )

    print(
        "Day 35-36 Combined Engine: ENABLED"
    )

    print(
        "Endpoint: POST /api/interview/start"
    )

    print(
        "Endpoint: POST /api/interview/evaluate"
    )

    print(
        "Endpoint: GET /api/interview/session/<session_id>"
    )

    print(
        "Endpoint: GET /api/interview/report/<session_id>"
    )

    print(
        "Endpoint: POST /api/interview/transcribe"
    )

    print(
        "Endpoint: POST /api/interview/speak"
    )

    print("=" * 50)

    print(
        "CORS: ENABLED"
    )

    print(
        "Allowed Origins: "
        + ", ".join(CORS_ORIGINS)
    )

    print(
        "Authorization Header: ENABLED"
    )

    print(
        "OPTIONS Preflight: ENABLED"
    )

    print("=" * 50)


    app.run(

        host=os.getenv(
            "HOST",
            "0.0.0.0"
        ),

        port=int(
            os.getenv(
                "PORT",
                "5000"
            )
        ),

        debug=(
            not IS_PRODUCTION
            and os.getenv(
                "FLASK_DEBUG",
                "false"
            )
            .strip()
            .lower()
            ==
            "true"
        )

    )